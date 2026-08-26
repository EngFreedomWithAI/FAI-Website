#!/usr/bin/env python3
"""Generate the Freedom With AI LLC Mutual NDA Cover Page as fillable PDF and DOCX."""

from __future__ import annotations

import io
from pathlib import Path

from PIL import Image
from pypdf import PdfReader, PdfWriter
from pypdf.generic import (
    ArrayObject,
    BooleanObject,
    DictionaryObject,
    FloatObject,
    NameObject,
    NumberObject,
    TextStringObject,
)
from reportlab.lib.colors import Color, HexColor, white
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas
from reportlab.platypus import Paragraph

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
LOGO_SRC = ROOT / "Assets" / "Logo1 - #3B3964.png"
OUT_DIR = Path(__file__).resolve().parent / "templates"
PDF_PATH = OUT_DIR / "Freedom-With-AI-Mutual-NDA-Cover-Page.pdf"
DOCX_PATH = OUT_DIR / "Freedom-With-AI-Mutual-NDA-Cover-Page.docx"

STANDARD_TERMS_URL = "https://commonpaper.com/standards/mutual-nda/1.0/"
STANDARD_TERMS_LABEL = "Common Paper Mutual NDA Standard Terms v1.0"

SLATE = HexColor("#3B3964")
BLUE = HexColor("#5165A2")
INK = HexColor("#1C1B1A")
MUTED = HexColor("#6B6973")
RULE = HexColor("#D8D4CB")
FIELD_BG = HexColor("#FFFFFF")
FIELD_BORDER = HexColor("#C8C4B8")
CARD_BG = HexColor("#FBF9F4")
QUOTE_BG = HexColor("#F4F1E8")
PAGE_BG = HexColor("#FFFcf7")

SLATE_RGB = RGBColor(0x3B, 0x39, 0x64)
BLUE_RGB = RGBColor(0x51, 0x65, 0xA2)
INK_RGB = RGBColor(0x1C, 0x1B, 0x1A)
MUTED_RGB = RGBColor(0x6B, 0x69, 0x73)
RULE_HEX = "D8D4CB"
FIELD_BG_HEX = "F7F5EF"
CARD_BG_HEX = "FBF9F4"
QUOTE_BG_HEX = "F4F1E8"

PURPOSE = (
    "Evaluating, discussing, negotiating, and pursuing a potential business "
    "relationship between the parties, including scoping potential services and "
    "related technical, commercial, product, and security discussions."
)
CHANGE_SENTENCE = (
    "Confidential Information includes any such information disclosed or made "
    "available before or after the Effective Date."
)

PAGE_W, PAGE_H = letter
MARGIN_X = 0.7 * inch
CONTENT_W = PAGE_W - 2 * MARGIN_X


def register_fonts() -> None:
    ttc = "/System/Library/Fonts/HelveticaNeue.ttc"
    pdfmetrics.registerFont(TTFont("HN", ttc, subfontIndex=0))
    pdfmetrics.registerFont(TTFont("HN-Bold", ttc, subfontIndex=1))
    pdfmetrics.registerFont(TTFont("HN-Med", ttc, subfontIndex=10))


def transparent_logo() -> io.BytesIO:
    im = Image.open(LOGO_SRC).convert("RGBA")
    pixels = []
    for r, g, b, a in im.get_flattened_data():
        if r > 248 and g > 248 and b > 248:
            pixels.append((255, 255, 255, 0))
        else:
            pixels.append((r, g, b, a))
    im.putdata(pixels)
    buf = io.BytesIO()
    im.save(buf, format="PNG")
    buf.seek(0)
    return buf


def paragraph(text: str, size: float, color: Color, align=TA_LEFT, leading=None, font="HN"):
    style = ParagraphStyle(
        name=f"p-{size}-{align}-{font}",
        fontName=font,
        fontSize=size,
        leading=leading or size * 1.35,
        textColor=color,
        alignment=align,
    )
    return Paragraph(text, style)


def draw_para(c: canvas.Canvas, text: str, x: float, y_top: float, width: float, **kwargs) -> float:
    p = paragraph(text, **kwargs)
    w, h = p.wrap(width, 400)
    p.drawOn(c, x, y_top - h)
    return h


def draw_round_rect(c: canvas.Canvas, x, y, w, h, radius, fill, stroke=None, sw=0.6):
    c.saveState()
    c.setFillColor(fill)
    if stroke:
        c.setStrokeColor(stroke)
        c.setLineWidth(sw)
        c.roundRect(x, y, w, h, radius, fill=1, stroke=1)
    else:
        c.setStrokeColor(fill)
        c.roundRect(x, y, w, h, radius, fill=1, stroke=0)
    c.restoreState()


def add_text_field(
    c: canvas.Canvas,
    name: str,
    x: float,
    y: float,
    w: float,
    h: float,
    value: str = "",
    tooltip: str = "",
    maxlen: int = 200,
) -> None:
    c.acroForm.textfield(
        name=name,
        tooltip=tooltip or name,
        x=x,
        y=y,
        width=w,
        height=h,
        value=value,
        fontName="Helvetica",
        fontSize=8.5,
        textColor=INK,
        fillColor=FIELD_BG,
        borderColor=FIELD_BORDER,
        borderWidth=0.6,
        forceBorder=True,
        fieldFlags="",
        maxlen=maxlen,
    )


def labeled_field(
    c: canvas.Canvas,
    label: str,
    name: str,
    x: float,
    y_top: float,
    width: float,
    value: str = "",
    tooltip: str = "",
    field_h: float = 16,
) -> float:
    c.setFillColor(MUTED)
    c.setFont("HN-Med", 7)
    c.drawString(x, y_top - 8, label.upper())
    field_y = y_top - 8 - 4 - field_h
    add_text_field(c, name, x, field_y, width, field_h, value=value, tooltip=tooltip or label)
    return 8 + 4 + field_h + 8


def add_signature_fields(pdf_path: Path, signatures: list[dict]) -> None:
    reader = PdfReader(str(pdf_path))
    writer = PdfWriter()
    writer.append(reader)

    page = writer.pages[0]
    acro = writer.root_object["/AcroForm"].get_object()
    acro[NameObject("/NeedAppearances")] = BooleanObject(True)
    acro[NameObject("/SigFlags")] = NumberObject(3)
    fields = acro["/Fields"]

    for sig in signatures:
        annot = DictionaryObject()
        annot.update(
            {
                NameObject("/Type"): NameObject("/Annot"),
                NameObject("/Subtype"): NameObject("/Widget"),
                NameObject("/FT"): NameObject("/Sig"),
                NameObject("/T"): TextStringObject(sig["name"]),
                NameObject("/TU"): TextStringObject(sig["tooltip"]),
                NameObject("/Rect"): ArrayObject([FloatObject(v) for v in sig["rect"]]),
                NameObject("/F"): NumberObject(4),
                NameObject("/P"): page.indirect_reference,
                NameObject("/MK"): DictionaryObject(
                    {
                        NameObject("/BC"): ArrayObject(
                            [FloatObject(0.83), FloatObject(0.82), FloatObject(0.78)]
                        ),
                        NameObject("/BG"): ArrayObject(
                            [FloatObject(0.969), FloatObject(0.961), FloatObject(0.937)]
                        ),
                    }
                ),
            }
        )
        ref = writer._add_object(annot)
        if "/Annots" in page:
            page[NameObject("/Annots")].append(ref)
        else:
            page[NameObject("/Annots")] = ArrayObject([ref])
        fields.append(ref)

    with pdf_path.open("wb") as f:
        writer.write(f)


def build_pdf(logo: io.BytesIO) -> None:
    register_fonts()
    c = canvas.Canvas(str(PDF_PATH), pagesize=letter)
    c.setTitle("Freedom With AI LLC - Mutual NDA Cover Page")
    c.setAuthor("Freedom With AI LLC")
    c.setSubject("Mutual NDA Cover Page incorporating Common Paper Mutual NDA Standard Terms v1.0")
    c.setCreator("Freedom With AI LLC")

    c.setFillColor(PAGE_BG)
    c.rect(0, 0, PAGE_W, PAGE_H, fill=1, stroke=0)
    c.setFillColor(SLATE)
    c.rect(0, PAGE_H - 6, PAGE_W, 6, fill=1, stroke=0)

    y = PAGE_H - 28
    logo_size = 46
    c.drawImage(
        ImageReader(logo),
        (PAGE_W - logo_size) / 2,
        y - logo_size,
        width=logo_size,
        height=logo_size,
        mask="auto",
        preserveAspectRatio=True,
        anchor="c",
    )
    y -= logo_size + 10

    c.setFillColor(SLATE)
    c.setFont("HN-Med", 8)
    c.drawCentredString(PAGE_W / 2, y, "FREEDOM WITH AI LLC")
    y -= 16
    c.setFont("HN-Bold", 16)
    c.drawCentredString(PAGE_W / 2, y, "Mutual NDA Cover Page")
    y -= 13
    c.setFillColor(MUTED)
    c.setFont("HN", 8)
    c.drawCentredString(PAGE_W / 2, y, f"For use with {STANDARD_TERMS_LABEL}")
    url_w = c.stringWidth(STANDARD_TERMS_URL, "HN", 7.5)
    y -= 11
    c.setFillColor(BLUE)
    c.setFont("HN", 7.5)
    c.drawCentredString(PAGE_W / 2, y, STANDARD_TERMS_URL)
    c.linkURL(
        STANDARD_TERMS_URL,
        (PAGE_W / 2 - url_w / 2 - 2, y - 2, PAGE_W / 2 + url_w / 2 + 2, y + 9),
        relative=0,
        thickness=0,
    )
    y -= 12
    c.setStrokeColor(BLUE)
    c.setLineWidth(0.8)
    c.line(MARGIN_X, y, PAGE_W - MARGIN_X, y)
    y -= 16

    h = draw_para(
        c,
        "This Mutual NDA Cover Page is entered into between the following parties.",
        MARGIN_X,
        y,
        CONTENT_W,
        size=9,
        color=INK,
        font="HN",
    )
    y -= h + 12

    gap = 14
    card_w = (CONTENT_W - gap) / 2
    card_h = 238
    left_x = MARGIN_X
    right_x = MARGIN_X + card_w + gap
    card_bottom = y - card_h

    signatures = []
    draw_party_card(
        c,
        x=left_x,
        bottom=card_bottom,
        width=card_w,
        height=card_h,
        eyebrow="Disclosing / Receiving Party 1",
        legal_name="Freedom With AI LLC",
        prefix="party1",
        title_value="Co-Founder, Authorized Signatory",
        signatures=signatures,
    )
    draw_party_card(
        c,
        x=right_x,
        bottom=card_bottom,
        width=card_w,
        height=card_h,
        eyebrow="Disclosing / Receiving Party 2",
        legal_name="",
        prefix="party2",
        title_value="",
        signatures=signatures,
    )
    y = card_bottom - 14

    y = draw_section_label(c, "Covered Agreement", MARGIN_X, y)
    h = draw_para(
        c,
        f"The parties agree that this Cover Page is governed by the <b>{STANDARD_TERMS_LABEL}</b>, "
        f"which are incorporated by reference, subject to the changes below.",
        MARGIN_X,
        y,
        CONTENT_W,
        size=8.5,
        color=INK,
        font="HN",
        leading=11.5,
    )
    y -= h + 12

    y = draw_section_label(c, "Key Terms", MARGIN_X, y)
    terms = [
        ("Purpose", PURPOSE),
        ("Effective Date", "The date of the last signature on this Cover Page."),
        ("MNDA Term", "2 years from the Effective Date."),
        (
            "Term of Confidentiality",
            "2 years from the date of last disclosure, except that trade secrets will remain "
            "protected for as long as they qualify as trade secrets under applicable law.",
        ),
        ("Governing Law", "California."),
        ("Jurisdiction", "State or federal courts located in California."),
    ]
    label_w = 122
    for i, (label, value) in enumerate(terms):
        p = paragraph(value, size=8.2, color=INK, font="HN", leading=11)
        vw, vh = p.wrap(CONTENT_W - label_w - 8, 200)
        row_h = max(vh, 10) + 7
        if i % 2 == 0:
            c.setFillColor(HexColor("#F7F5EF"))
            c.rect(MARGIN_X - 4, y - row_h + 4, CONTENT_W + 8, row_h, fill=1, stroke=0)
        c.setFillColor(SLATE)
        c.setFont("HN-Med", 7.5)
        c.drawString(MARGIN_X, y - 8, label.upper())
        p.drawOn(c, MARGIN_X + label_w, y - vh)
        y -= row_h
    y -= 8

    y = draw_section_label(c, "Changes to Standard Terms", MARGIN_X, y)
    h = draw_para(
        c,
        "Add the following sentence after the first sentence of Section 1:",
        MARGIN_X,
        y,
        CONTENT_W,
        size=8.5,
        color=INK,
        font="HN",
    )
    y -= h + 6
    quote_h = 28
    draw_round_rect(c, MARGIN_X, y - quote_h, CONTENT_W, quote_h, 3, QUOTE_BG)
    h = draw_para(
        c,
        f'"{CHANGE_SENTENCE}"',
        MARGIN_X + 10,
        y - 6,
        CONTENT_W - 20,
        size=8.2,
        color=SLATE,
        font="HN",
        leading=11,
    )
    y -= quote_h + 8
    c.setFillColor(INK)
    c.setFont("HN-Med", 8.5)
    c.drawString(MARGIN_X, y, "No other changes.")

    c.setFillColor(SLATE)
    c.rect(0, 0, PAGE_W, 28, fill=1, stroke=0)
    c.setFillColor(white)
    c.setFont("HN", 7)
    c.drawCentredString(
        PAGE_W / 2,
        11,
        "Freedom With AI LLC   ·   Mutual NDA Cover Page   ·   Incorporating Common Paper Mutual NDA Standard Terms v1.0",
    )

    c.save()
    add_signature_fields(PDF_PATH, signatures)


def draw_section_label(c: canvas.Canvas, text: str, x: float, y: float) -> float:
    c.setFillColor(SLATE)
    c.setFont("HN-Bold", 9)
    c.drawString(x, y - 2, text)
    return y - 14


def draw_party_card(
    c: canvas.Canvas,
    x: float,
    bottom: float,
    width: float,
    height: float,
    eyebrow: str,
    legal_name: str,
    prefix: str,
    title_value: str,
    signatures: list[dict],
) -> None:
    draw_round_rect(c, x, bottom, width, height, 5, CARD_BG, RULE, 0.7)
    c.setFillColor(SLATE)
    c.roundRect(x, bottom + height - 22, width, 22, 5, fill=1, stroke=0)
    c.rect(x, bottom + height - 22, width, 8, fill=1, stroke=0)
    c.setFillColor(white)
    c.setFont("HN-Med", 7.5)
    c.drawString(x + 10, bottom + height - 14, eyebrow.upper())

    pad = 10
    inner_x = x + pad
    inner_w = width - 2 * pad
    cursor = bottom + height - 22 - 10

    used = labeled_field(
        c,
        "Legal Name",
        f"{prefix}_legal_name",
        inner_x,
        cursor,
        inner_w,
        value=legal_name,
        tooltip="Exact legal entity name",
    )
    cursor -= used
    used = labeled_field(
        c,
        "Notice Email",
        f"{prefix}_notice_email",
        inner_x,
        cursor,
        inner_w,
        tooltip="Monitored notice email",
    )
    cursor -= used
    used = labeled_field(
        c,
        "Signature Name",
        f"{prefix}_signature_name",
        inner_x,
        cursor,
        inner_w,
        tooltip="Name of authorized signer",
    )
    cursor -= used
    used = labeled_field(
        c,
        "Title",
        f"{prefix}_title",
        inner_x,
        cursor,
        inner_w,
        value=title_value,
        tooltip="Signer title",
    )
    cursor -= used

    sig_w = inner_w * 0.64
    date_w = inner_w - sig_w - 8
    c.setFillColor(MUTED)
    c.setFont("HN-Med", 7)
    c.drawString(inner_x, cursor - 8, "SIGNATURE")
    c.drawString(inner_x + sig_w + 8, cursor - 8, "DATE")
    sig_h = 36
    sig_y = cursor - 8 - 4 - sig_h
    draw_round_rect(c, inner_x, sig_y, sig_w, sig_h, 2, FIELD_BG, FIELD_BORDER, 0.6)
    c.setStrokeColor(FIELD_BORDER)
    c.setLineWidth(0.6)
    c.line(inner_x + 8, sig_y + 10, inner_x + sig_w - 8, sig_y + 10)
    add_text_field(
        c,
        f"{prefix}_date",
        inner_x + sig_w + 8,
        sig_y,
        date_w,
        sig_h,
        tooltip="Signature date",
        maxlen=40,
    )
    signatures.append(
        {
            "name": f"{prefix}_signature",
            "tooltip": f"Click to sign as {eyebrow}",
            "rect": [inner_x, sig_y, inner_x + sig_w, sig_y + sig_h],
        }
    )


def set_run_font(run, name: str, size: float, color: RGBColor, bold: bool = False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.bold = bold


def set_cell_shading(cell, hex_color: str) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_margins(cell, top=40, bottom=40, left=80, right=80) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for edge, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)


def set_table_borders(table, color=RULE_HEX, sz="4") -> None:
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else OxmlElement("w:tblPr")
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), sz)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tblPr.append(borders)


def remove_table_borders(table) -> None:
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else OxmlElement("w:tblPr")
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tblPr.append(borders)


def add_bottom_border(paragraph, color="3B3964", sz="12") -> None:
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), sz)
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_content_control(paragraph, tag: str, alias: str, default: str = "") -> None:
    sdt = OxmlElement("w:sdt")
    sdtPr = OxmlElement("w:sdtPr")

    alias_el = OxmlElement("w:alias")
    alias_el.set(qn("w:val"), alias)
    sdtPr.append(alias_el)

    tag_el = OxmlElement("w:tag")
    tag_el.set(qn("w:val"), tag)
    sdtPr.append(tag_el)

    text_el = OxmlElement("w:text")
    sdtPr.append(text_el)

    if not default:
        sdtPr.append(OxmlElement("w:showingPlcHdr"))
        placeholder = OxmlElement("w:placeholder")
        doc_part = OxmlElement("w:docPart")
        doc_part.set(qn("w:val"), "DefaultPlaceholder_-1854013437")
        placeholder.append(doc_part)
        sdtPr.append(placeholder)

    sdt.append(sdtPr)
    sdt_content = OxmlElement("w:sdtContent")
    r = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Calibri")
    fonts.set(qn("w:hAnsi"), "Calibri")
    rPr.append(fonts)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "18")
    rPr.append(sz)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1C1B1A" if default else "9A9890")
    rPr.append(color)
    r.append(rPr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = default or alias
    r.append(t)
    sdt_content.append(r)
    sdt.append(sdt_content)
    paragraph._p.append(sdt)


def shade_paragraph(paragraph, hex_color: str) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    pPr.append(shd)


def tight(paragraph, before=0, after=0, line=None):
    pf = paragraph.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if line:
        pf.line_spacing = line


def add_centered(doc, text, size, color, bold=False, before=0, after=0, name="Calibri"):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tight(p, before, after)
    run = p.add_run(text)
    set_run_font(run, name, size, color, bold)
    return p


def add_body(doc, text, size=10, color=INK_RGB, bold=False, before=0, after=6):
    p = doc.add_paragraph()
    tight(p, before, after)
    run = p.add_run(text)
    set_run_font(run, "Calibri", size, color, bold)
    return p


def add_field_row(cell, label: str, tag: str, alias: str, default: str = "") -> None:
    lp = cell.add_paragraph()
    tight(lp, 2, 0)
    run = lp.add_run(label.upper())
    set_run_font(run, "Calibri", 7.5, MUTED_RGB, True)

    fp = cell.add_paragraph()
    tight(fp, 0, 1)
    shade_paragraph(fp, FIELD_BG_HEX)
    add_bottom_border(fp, RULE_HEX, "6")
    add_content_control(fp, tag, alias, default)


def build_docx(logo: io.BytesIO) -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.55)

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run(
        "Freedom With AI LLC   ·   Mutual NDA Cover Page   ·   Incorporating Common Paper Mutual NDA Standard Terms v1.0"
    )
    set_run_font(run, "Calibri", 8, SLATE_RGB)

    # Top brand rule via a full-width paragraph border
    brand = doc.add_paragraph()
    tight(brand, 0, 4)
    add_bottom_border(brand, "3B3964", "24")

    logo_p = doc.add_paragraph()
    logo_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tight(logo_p, 0, 4)
    run = logo_p.add_run()
    run.add_picture(logo, width=Inches(0.52))

    add_centered(doc, "FREEDOM WITH AI LLC", 8.5, SLATE_RGB, True, 1, 0)
    add_centered(doc, "Mutual NDA Cover Page", 16, SLATE_RGB, True, 1, 0)
    add_centered(doc, f"For use with {STANDARD_TERMS_LABEL}", 8.5, MUTED_RGB, False, 1, 0)
    url_p = add_centered(doc, STANDARD_TERMS_URL, 8, BLUE_RGB, False, 0, 4)
    add_bottom_border(url_p, "5165A2", "12")

    add_body(
        doc,
        "This Mutual NDA Cover Page is entered into between the following parties.",
        10,
        INK_RGB,
        before=6,
        after=6,
    )

    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_table_borders(table, RULE_HEX, "6")

    parties = [
        {
            "cell": table.cell(0, 0),
            "eyebrow": "Disclosing / Receiving Party 1",
            "prefix": "party1",
            "legal": "Freedom With AI LLC",
            "title": "Co-Founder, Authorized Signatory",
            "legal_alias": "Freedom With AI LLC",
            "name_alias": "Co-founder name",
            "email_alias": "Notice email",
            "title_alias": "Co-Founder, Authorized Signatory",
        },
        {
            "cell": table.cell(0, 1),
            "eyebrow": "Disclosing / Receiving Party 2",
            "prefix": "party2",
            "legal": "",
            "title": "",
            "legal_alias": "Client legal entity name",
            "name_alias": "Client signer name",
            "email_alias": "Client notice email",
            "title_alias": "Signer title",
        },
    ]

    for party in parties:
        cell = party["cell"]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        set_cell_shading(cell, CARD_BG_HEX)
        set_cell_margins(cell, 40, 50, 70, 70)

        # Clear default empty paragraph content by using the first paragraph as header
        header = cell.paragraphs[0]
        tight(header, 0, 6)
        shade_paragraph(header, "3B3964")
        header.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = header.add_run(party["eyebrow"].upper())
        set_run_font(run, "Calibri", 8, RGBColor(255, 255, 255), True)

        add_field_row(cell, "Legal Name", f"{party['prefix']}_legal_name", party["legal_alias"], party["legal"])
        add_field_row(cell, "Notice Email", f"{party['prefix']}_notice_email", party["email_alias"])
        add_field_row(cell, "Signature Name", f"{party['prefix']}_signature_name", party["name_alias"])
        add_field_row(cell, "Title", f"{party['prefix']}_title", party["title_alias"], party["title"])

        sig_label = cell.add_paragraph()
        tight(sig_label, 8, 0)
        run = sig_label.add_run("SIGNATURE")
        set_run_font(run, "Calibri", 8, MUTED_RGB, True)
        run2 = sig_label.add_run("                              DATE")
        set_run_font(run2, "Calibri", 8, MUTED_RGB, True)

        sig_line = cell.add_paragraph()
        tight(sig_line, 2, 2)
        add_content_control(sig_line, f"{party['prefix']}_signature", "Sign here")
        run = sig_line.add_run("          ")
        set_run_font(run, "Calibri", 11, INK_RGB)
        add_content_control(sig_line, f"{party['prefix']}_date", "Date")
        add_bottom_border(sig_line, RULE_HEX, "8")

    add_body(doc, "Covered Agreement", 11, SLATE_RGB, True, 8, 2)
    add_body(
        doc,
        f"The parties agree that this Cover Page is governed by the {STANDARD_TERMS_LABEL}, "
        "which are incorporated by reference, subject to the changes below.",
        10,
        INK_RGB,
        after=4,
    )

    add_body(doc, "Key Terms", 11, SLATE_RGB, True, 2, 2)
    terms = [
        ("Purpose", PURPOSE),
        ("Effective Date", "The date of the last signature on this Cover Page."),
        ("MNDA Term", "2 years from the Effective Date."),
        (
            "Term of Confidentiality",
            "2 years from the date of last disclosure, except that trade secrets will remain "
            "protected for as long as they qualify as trade secrets under applicable law.",
        ),
        ("Governing Law", "California."),
        ("Jurisdiction", "State or federal courts located in California."),
    ]
    terms_table = doc.add_table(rows=len(terms), cols=2)
    set_table_borders(terms_table, "FFFFFF", "0")
    remove_table_borders(terms_table)
    for i, (label, value) in enumerate(terms):
        left, right = terms_table.cell(i, 0), terms_table.cell(i, 1)
        set_cell_margins(left, 28, 28, 50, 30)
        set_cell_margins(right, 28, 28, 30, 50)
        if i % 2 == 0:
            set_cell_shading(left, FIELD_BG_HEX)
            set_cell_shading(right, FIELD_BG_HEX)
        lp = left.paragraphs[0]
        tight(lp, 0, 0)
        run = lp.add_run(label.upper())
        set_run_font(run, "Calibri", 8.5, SLATE_RGB, True)
        rp = right.paragraphs[0]
        tight(rp, 0, 0)
        run = rp.add_run(value)
        set_run_font(run, "Calibri", 10, INK_RGB)
    for cell in terms_table.columns[0].cells:
        cell.width = Inches(1.7)
    for cell in terms_table.columns[1].cells:
        cell.width = Inches(5.4)

    add_body(doc, "Changes to Standard Terms", 11, SLATE_RGB, True, 8, 2)
    add_body(doc, "Add the following sentence after the first sentence of Section 1:", 10, INK_RGB, after=2)
    quote = add_body(doc, f'"{CHANGE_SENTENCE}"', 10, SLATE_RGB, after=2)
    shade_paragraph(quote, QUOTE_BG_HEX)
    add_body(doc, "No other changes.", 10, INK_RGB, True, 2, 0)

    doc.save(str(DOCX_PATH))


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    logo = transparent_logo()
    build_pdf(io.BytesIO(logo.getvalue()))
    build_docx(io.BytesIO(logo.getvalue()))
    print(f"Wrote {PDF_PATH}")
    print(f"Wrote {DOCX_PATH}")


if __name__ == "__main__":
    main()
